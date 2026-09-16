import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_callout(doc, text, title="LƯU Ý QUAN TRỌNG / HỒ SƠ QUYẾT ĐỊNH (ADR)"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "EBF3FA")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=150)
    
    # Left border thick navy, other borders none
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="1F4E79"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"📌 {title}\n")
    run_t.bold = True
    run_t.font.name = "Arial"
    run_t.font.size = Pt(10.5)
    run_t.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    
    run_b = p.add_run(text)
    run_b.font.name = "Arial"
    run_b.font.size = Pt(10)
    run_b.font.color.rgb = RGBColor(0x2A, 0x2A, 0x2A)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def style_table(table, col_widths, headers, data):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Header row
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].width = Inches(col_widths[i])
        set_cell_background(hdr_cells[i], "1F4E79")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(h)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    # Data rows
    for r_idx, row_data in enumerate(data):
        row = table.add_row()
        fill = "F9FAFB" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.width = Inches(col_widths[c_idx])
            set_cell_background(cell, fill)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            if c_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(val) <= 12 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(val)
            run.font.name = "Arial"
            run.font.size = Pt(9)
            if "LỰA CHỌN TỐI ƯU" in val or "TỔNG ĐIỂM" in val:
                run.bold = True
                run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
            else:
                run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

def add_heading_with_spacing(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.keep_with_next = True
    if level == 1:
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(8)
        for r in h.runs:
            r.font.name = "Arial"
            r.font.size = Pt(16)
            r.bold = True
            r.font.color.rgb = RGBColor(0x0F, 0x2A, 0x4A)
    elif level == 2:
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
        for r in h.runs:
            r.font.name = "Arial"
            r.font.size = Pt(13)
            r.bold = True
            r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    elif level == 3:
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
        for r in h.runs:
            r.font.name = "Arial"
            r.font.size = Pt(11)
            r.bold = True
            r.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
    return h

def add_body_p(doc, text, bold_prefix=None, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = "Arial"
        r_pre.font.size = Pt(10.5)
        r_pre.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    return p

def add_bullet_p(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = "Arial"
        r_pre.font.size = Pt(10.5)
        r_pre.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    return p

def add_image_safe(doc, img_path, caption, width_in=5.8):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(3)
        run_img = p_img.add_run()
        try:
            run_img.add_picture(img_path, width=Inches(width_in))
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(0)
            p_cap.paragraph_format.space_after = Pt(10)
            run_cap = p_cap.add_run(f"Hình minh họa: {caption}")
            run_cap.italic = True
            run_cap.font.name = "Arial"
            run_cap.font.size = Pt(9.5)
            run_cap.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        except Exception as e:
            print(f"Failed adding picture {img_path}: {e}")

print("Docx generator helper functions loaded.")
